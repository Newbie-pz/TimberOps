"""XLSX weighing-history export integration tests."""

from datetime import datetime, timedelta, timezone
from io import BytesIO
from uuid import UUID

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine
from sqlalchemy.orm import Session

from app.models.weighing import WeighingTask
from app.services.export_service import EXPORT_HEADERS


def _vehicle(client: TestClient, plate_number: str) -> dict[str, object]:
    response = client.post(
        "/api/v1/vehicles",
        json={
            "plate_number": plate_number,
            "driver_name": "张师傅",
            "vehicle_type": "LARGE",
            "allowed_gross_weight_tons": "49.000",
        },
    )
    assert response.status_code == 201
    return response.json()


def test_export_returns_xlsx_and_applies_all_filters(
    api_client: TestClient,
    db_engine: Engine,
) -> None:
    customer_response = api_client.post(
        "/api/v1/customers",
        json={"name": "二连浩特口岸客户"},
    )
    customer = customer_response.json()
    timber_vehicle = _vehicle(api_client, "蒙H70001")
    coal_vehicle = _vehicle(api_client, "蒙H70002")
    timber_response = api_client.post(
        "/api/v1/weighing/tasks",
        json={
            "vehicle_id": timber_vehicle["id"],
            "customer_id": customer["id"],
            "cargo_type": "TIMBER",
            "cargo_name": "落叶松原木",
            "cargo_remark": "俄罗斯进口",
        },
    )
    coal_response = api_client.post(
        "/api/v1/weighing/tasks",
        json={
            "vehicle_id": coal_vehicle["id"],
            "cargo_type": "COAL",
            "cargo_name": "原煤",
        },
    )
    assert timber_response.status_code == 201
    assert coal_response.status_code == 201
    timber_task_id = timber_response.json()["id"]
    assert api_client.post(
        f"/api/v1/weighing/tasks/{timber_task_id}/tare",
        json={"weight_tons": "15.820"},
    ).status_code == 200
    assert api_client.post(
        f"/api/v1/weighing/tasks/{timber_task_id}/wait-gross"
    ).status_code == 200
    assert api_client.post(
        f"/api/v1/weighing/tasks/{timber_task_id}/gross",
        json={"weight_tons": "47.000"},
    ).status_code == 200
    assert api_client.post(
        f"/api/v1/weighing/tasks/{timber_task_id}/complete"
    ).status_code == 200

    with Session(db_engine) as session:
        old_task = session.get(WeighingTask, UUID(coal_response.json()["id"]))
        assert old_task is not None
        old_task.created_at = datetime(2020, 1, 1, tzinfo=timezone.utc)
        session.commit()

    today = datetime.now(timezone(timedelta(hours=8))).date().isoformat()
    response = api_client.get(
        "/api/v1/export/weighing",
        params={
            "start_date": today,
            "end_date": today,
            "vehicle_id": timber_vehicle["id"],
            "customer_id": customer["id"],
            "cargo_type": "TIMBER",
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.content.startswith(b"PK")
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    worksheet = workbook["称重历史"]
    rows = list(worksheet.iter_rows(values_only=True))
    assert rows[0] == EXPORT_HEADERS
    assert len(rows) == 2
    assert rows[1][2] == "蒙H70001"
    assert rows[1][4] == "二连浩特口岸客户"
    assert rows[1][5:8] == ("TIMBER", "落叶松原木", "俄罗斯进口")
    assert rows[1][8:11] == (15.82, 47, 31.18)
    assert rows[1][11:13] == ("COMPLETED", "NORMAL")


def test_export_without_dates_returns_full_history(api_client: TestClient) -> None:
    first = _vehicle(api_client, "蒙H71001")
    second = _vehicle(api_client, "蒙H71002")
    for vehicle, cargo_type in ((first, "ORE"), (second, "COAL")):
        response = api_client.post(
            "/api/v1/weighing/tasks",
            json={"vehicle_id": vehicle["id"], "cargo_type": cargo_type},
        )
        assert response.status_code == 201

    response = api_client.get("/api/v1/export/weighing")
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    rows = list(workbook["称重历史"].iter_rows(values_only=True))

    assert response.status_code == 200
    assert len(rows) == 3


def test_export_rejects_reversed_date_range(api_client: TestClient) -> None:
    response = api_client.get(
        "/api/v1/export/weighing",
        params={"start_date": "2026-09-16", "end_date": "2026-09-15"},
    )

    assert response.status_code == 422
