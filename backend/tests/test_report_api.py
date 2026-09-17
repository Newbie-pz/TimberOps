"""Report API contracts, permissions, and XLSX generation."""

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine, select
from sqlalchemy.orm import Session

from app.models.rbac import Role, UserRole
from app.models.user import User
from app.security.jwt import create_access_token


def _headers_for_role(engine: Engine, role_name: str) -> dict[str, str]:
    with Session(engine) as session:
        role = session.scalar(select(Role).where(Role.name == role_name))
        assert role is not None
        username = f"report_{role_name.lower()}"
        user = User(username=username, password_hash="not-used", real_name=username)
        session.add(user)
        session.flush()
        session.add(UserRole(user_id=user.id, role_id=role.id))
        session.commit()
        token = create_access_token(user_id=user.id, username=user.username)
    return {"Authorization": f"Bearer {token}"}


def test_daily_and_monthly_report_api_return_structured_empty_data(
    api_client: TestClient,
) -> None:
    daily = api_client.get(
        "/api/v1/reports/daily",
        params={"report_date": "2026-09-17"},
    )
    monthly = api_client.get(
        "/api/v1/reports/monthly",
        params={"year": 2026, "month": 9},
    )

    assert daily.status_code == 200
    assert daily.json()["report_type"] == "daily"
    assert daily.json()["period_start"] == "2026-09-17"
    assert daily.json()["completed_net_weight_tons"] == "0.000"
    assert monthly.status_code == 200
    assert monthly.json()["report_type"] == "monthly"
    assert monthly.json()["period_end"] == "2026-09-30"
    assert monthly.json()["income"] == "0.00"


def test_report_export_generates_valid_xlsx(api_client: TestClient) -> None:
    response = api_client.get(
        "/api/v1/reports/export",
        params={
            "report_type": "daily",
            "report_date": "2026-09-17",
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert workbook.sheetnames == ["经营概览", "货物排行", "车辆排行"]
    assert workbook["经营概览"]["A1"].value == "指标"
    assert workbook["经营概览"]["B2"].value == "日报"


def test_viewer_can_view_but_cannot_export_reports(
    api_client: TestClient,
    db_engine: Engine,
) -> None:
    headers = _headers_for_role(db_engine, "VIEWER")

    viewed = api_client.get("/api/v1/reports/daily", headers=headers)
    exported = api_client.get(
        "/api/v1/reports/export",
        params={"report_type": "daily"},
        headers=headers,
    )

    assert viewed.status_code == 200
    assert exported.status_code == 403


def test_operator_can_view_and_export_reports(
    api_client: TestClient,
    db_engine: Engine,
) -> None:
    headers = _headers_for_role(db_engine, "OPERATOR")

    viewed = api_client.get("/api/v1/reports/monthly", headers=headers)
    exported = api_client.get(
        "/api/v1/reports/export",
        params={"report_type": "monthly"},
        headers=headers,
    )

    assert viewed.status_code == 200
    assert exported.status_code == 200
