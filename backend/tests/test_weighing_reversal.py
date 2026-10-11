"""Completed weighing reversal invariants, permissions, and query semantics."""

from datetime import datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.domain.enums import (
    CargoType,
    PaymentStatus,
    VehicleType,
    WeighingStatus,
    WeightResult,
    WeightSource,
    WeightType,
)
from app.domain.exceptions import (
    InvalidStateError,
    PaidWeighingRequiresRefundError,
    ReversedBillingImmutableError,
)
from app.models.audit_log import AuditLog
from app.models.billing import BillingRecord
from app.models.rbac import Role, UserRole
from app.models.user import User
from app.models.vehicle import Vehicle
from app.models.weighing import WeighingRecord, WeighingTask
from app.models.weighing_reversal import WeighingReversal
from app.schemas.weighing_reversal import WeighingReversalCreate
from app.security.jwt import create_access_token
from app.services.analytics_service import AnalyticsService
from app.services.billing_service import BillingService
from app.services.dashboard_service import DashboardService
from app.services.export_service import ExportService
from app.services.report_service import ReportService
from app.services.weighing_reversal_service import WeighingReversalService


def _role_headers(engine: Engine, role_name: str) -> dict[str, str]:
    with Session(engine) as session:
        role = session.scalar(select(Role).where(Role.name == role_name))
        assert role is not None
        user = User(
            username=f"reverse_{role_name.lower()}",
            password_hash="not-used",
            real_name=f"{role_name} 冲正测试",
        )
        session.add(user)
        session.flush()
        session.add(UserRole(user_id=user.id, role_id=role.id))
        session.commit()
        token = create_access_token(user_id=user.id, username=user.username)
    return {"Authorization": f"Bearer {token}"}


def _seed_task(
    session: Session,
    *,
    suffix: str,
    status: WeighingStatus = WeighingStatus.COMPLETED,
    payment_status: PaymentStatus = PaymentStatus.UNPAID,
) -> tuple[WeighingTask, BillingRecord | None]:
    vehicle = Vehicle(
        plate_number=f"蒙R-{suffix}",
        vehicle_type=VehicleType.MEDIUM,
        allowed_gross_weight_tons=Decimal("49.000"),
    )
    session.add(vehicle)
    session.flush()
    now = datetime.now(timezone.utc)
    task = WeighingTask(
        task_no=f"REV-{suffix}",
        vehicle_id=vehicle.id,
        cargo_type=CargoType.TIMBER,
        tare_weight_tons=(Decimal("15.000") if status is WeighingStatus.COMPLETED else None),
        gross_weight_tons=(Decimal("45.000") if status is WeighingStatus.COMPLETED else None),
        net_weight_tons=(Decimal("30.000") if status is WeighingStatus.COMPLETED else None),
        allowed_gross_weight_tons=Decimal("49.000"),
        overweight_tons=Decimal("0.000"),
        status=status,
        weight_result=(
            WeightResult.NORMAL
            if status is WeighingStatus.COMPLETED
            else WeightResult.PENDING
        ),
        completed_at=now if status is WeighingStatus.COMPLETED else None,
    )
    session.add(task)
    session.flush()
    billing = None
    if status is WeighingStatus.COMPLETED:
        session.add_all(
            [
                WeighingRecord(
                    weighing_task_id=task.id,
                    weight_type=WeightType.TARE,
                    weight_tons=Decimal("15.000"),
                    sequence_no=1,
                    recorded_at=now,
                    source=WeightSource.MANUAL,
                ),
                WeighingRecord(
                    weighing_task_id=task.id,
                    weight_type=WeightType.GROSS,
                    weight_tons=Decimal("45.000"),
                    sequence_no=2,
                    recorded_at=now,
                    source=WeightSource.MANUAL,
                ),
            ]
        )
        billing = BillingRecord(
            weighing_task_id=task.id,
            vehicle_id=vehicle.id,
            vehicle_type_snapshot=VehicleType.MEDIUM,
            fee_amount=Decimal("30.00"),
            payment_status=payment_status,
        )
        session.add(billing)
    session.commit()
    return task, billing


def test_only_completed_task_can_be_reversed(db_session: Session) -> None:
    task, _ = _seed_task(
        db_session,
        suffix="PENDING",
        status=WeighingStatus.WAIT_TARE,
    )
    operator = User(username="reverser", password_hash="x", real_name="管理员")
    db_session.add(operator)
    db_session.commit()

    with pytest.raises(InvalidStateError):
        WeighingReversalService(db_session).reverse(
            task.id,
            WeighingReversalCreate(reason="录入错误"),
            operator_id=operator.id,
        )


def test_reversal_is_idempotent_and_preserves_original_facts(
    db_session: Session,
) -> None:
    task, billing = _seed_task(db_session, suffix="SUCCESS")
    operator = User(username="admin_reverse", password_hash="x", real_name="管理员")
    db_session.add(operator)
    db_session.commit()
    original = (task.status, task.tare_weight_tons, task.gross_weight_tons, task.net_weight_tons)
    records_before = list(
        db_session.scalars(
            select(WeighingRecord).where(WeighingRecord.weighing_task_id == task.id)
        )
    )
    service = WeighingReversalService(db_session)

    first = service.reverse(
        task.id,
        WeighingReversalCreate(reason="毛重录入错误"),
        operator_id=operator.id,
    )
    second = service.reverse(
        task.id,
        WeighingReversalCreate(reason="重复请求不得覆盖原因"),
        operator_id=operator.id,
    )
    db_session.refresh(task)

    assert second.id == first.id
    assert second.reason == "毛重录入错误"
    assert (task.status, task.tare_weight_tons, task.gross_weight_tons, task.net_weight_tons) == original
    records_after = list(
        db_session.scalars(
            select(WeighingRecord).where(WeighingRecord.weighing_task_id == task.id)
        )
    )
    assert [(r.id, r.weight_tons) for r in records_after] == [
        (r.id, r.weight_tons) for r in records_before
    ]
    assert first.billing_record_id == billing.id  # type: ignore[union-attr]
    assert first.billing_payment_status_snapshot is PaymentStatus.UNPAID
    audits = list(
        db_session.scalars(
            select(AuditLog).where(AuditLog.action == "WEIGHING_REVERSED")
        )
    )
    assert len(audits) == 1
    assert audits[0].reason == "毛重录入错误"


def test_paid_task_requires_future_refund_flow(db_session: Session) -> None:
    task, _ = _seed_task(
        db_session,
        suffix="PAID",
        payment_status=PaymentStatus.PAID,
    )
    operator = User(username="paid_admin", password_hash="x", real_name="管理员")
    db_session.add(operator)
    db_session.commit()

    with pytest.raises(PaidWeighingRequiresRefundError):
        WeighingReversalService(db_session).reverse(
            task.id,
            WeighingReversalCreate(reason="发现错误"),
            operator_id=operator.id,
        )


def test_waived_billing_can_be_reversed_and_remains_waived(
    db_session: Session,
) -> None:
    task, billing = _seed_task(
        db_session,
        suffix="WAIVED",
        payment_status=PaymentStatus.WAIVED,
    )
    operator = User(username="waived_admin", password_hash="x", real_name="管理员")
    db_session.add(operator)
    db_session.commit()

    reversal = WeighingReversalService(db_session).reverse(
        task.id,
        WeighingReversalCreate(reason="免除费用后确认业务录入错误"),
        operator_id=operator.id,
    )
    assert billing is not None
    db_session.refresh(billing)

    assert reversal.billing_payment_status_snapshot is PaymentStatus.WAIVED
    assert billing.payment_status is PaymentStatus.WAIVED


def test_reversed_billing_cannot_be_paid(db_session: Session) -> None:
    task, billing = _seed_task(db_session, suffix="BILLING")
    operator = User(username="billing_admin", password_hash="x", real_name="管理员")
    db_session.add(operator)
    db_session.commit()
    WeighingReversalService(db_session).reverse(
        task.id,
        WeighingReversalCreate(reason="业务无效"),
        operator_id=operator.id,
    )

    with pytest.raises(ReversedBillingImmutableError):
        BillingService(db_session).mark_paid(
            billing.id,  # type: ignore[union-attr]
            operator_id=operator.id,
        )


def test_database_unique_constraint_prevents_two_reversals(db_session: Session) -> None:
    task, _ = _seed_task(db_session, suffix="UNIQUE")
    first = WeighingReversal(original_task_id=task.id, reason="first")
    db_session.add(first)
    db_session.commit()
    db_session.add(WeighingReversal(original_task_id=task.id, reason="second"))

    with pytest.raises(IntegrityError):
        db_session.commit()
    db_session.rollback()


@pytest.mark.parametrize("role_name", ["VIEWER", "OPERATOR"])
def test_non_admin_roles_cannot_reverse(
    api_client: TestClient,
    db_engine: Engine,
    role_name: str,
) -> None:
    with Session(db_engine) as session:
        task, _ = _seed_task(session, suffix=role_name)
        task_id = task.id

    response = api_client.post(
        f"/api/v1/weighing/tasks/{task_id}/reverse",
        headers=_role_headers(db_engine, role_name),
        json={"reason": "不应允许"},
    )

    assert response.status_code == 403


def test_admin_reverse_api_and_reason_validation(
    api_client: TestClient,
    db_engine: Engine,
) -> None:
    with Session(db_engine) as session:
        task, _ = _seed_task(session, suffix="API")
        task_id = task.id

    invalid = api_client.post(
        f"/api/v1/weighing/tasks/{task_id}/reverse",
        json={"reason": "   "},
    )
    success = api_client.post(
        f"/api/v1/weighing/tasks/{task_id}/reverse",
        json={"reason": "司机确认录入错误"},
    )
    repeated = api_client.post(
        f"/api/v1/weighing/tasks/{task_id}/reverse",
        json={"reason": "retry"},
    )

    assert invalid.status_code == 422
    assert success.status_code == 200
    assert repeated.status_code == 200
    assert repeated.json()["id"] == success.json()["id"]
    assert success.json()["operator_name"] == "测试管理员"


def test_operating_queries_exclude_reversal_but_history_exposes_it(
    db_session: Session,
) -> None:
    task, _ = _seed_task(db_session, suffix="QUERIES")
    operator = User(username="query_admin", password_hash="x", real_name="查询管理员")
    db_session.add(operator)
    db_session.commit()
    WeighingReversalService(db_session).reverse(
        task.id,
        WeighingReversalCreate(reason="统计冲正"),
        operator_id=operator.id,
    )

    dashboard = DashboardService(db_session).get_overview()
    local_date = datetime.now(timezone(timedelta(hours=8))).date()
    daily = ReportService(db_session).daily(local_date)
    monthly = ReportService(db_session).monthly(
        year=local_date.year,
        month=local_date.month,
    )
    analytics = AnalyticsService(db_session)
    summary = analytics.get_today_weighing_summary()
    detail = analytics.get_weighing_task_detail(task_no=task.task_no)
    history = analytics.get_vehicle_weighing_history(
        plate_number=task.vehicle.plate_number
    )
    workbook = load_workbook(
        BytesIO(ExportService(db_session).export_weighing()),
        data_only=True,
    )
    export_rows = list(workbook["称重历史"].iter_rows(values_only=True))

    assert dashboard.today_completed_task_count == 0
    assert dashboard.today_completed_net_weight_tons == Decimal("0.000")
    assert dashboard.today_income == Decimal("0.00")
    assert daily.completed_task_count == monthly.completed_task_count == 0
    assert daily.completed_net_weight_tons == monthly.completed_net_weight_tons == Decimal("0.000")
    assert summary["completed_tasks"] == 0
    assert detail["reversed"] is True
    assert detail["reversal_reason"] == "统计冲正"
    assert detail["reversed_by"] == "查询管理员"
    assert history["tasks"][0]["reversed"] is True
    assert export_rows[0][-3:] == ("是否冲正", "冲正时间", "冲正原因")
    assert export_rows[1][-3] == "是"
    assert export_rows[1][-1] == "统计冲正"
