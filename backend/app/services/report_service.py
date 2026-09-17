"""Read-only periodic operating reports and their XLSX representation."""

from calendar import monthrange
from collections.abc import Callable
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.db.base import utc_now
from app.domain.enums import CargoType, PaymentStatus, WeighingStatus
from app.domain.exceptions import BusinessRuleError
from app.models.billing import BillingRecord
from app.models.vehicle import Vehicle
from app.models.weighing import WeighingTask
from app.schemas.report import (
    BusinessReport,
    ReportCargoRankingItem,
    ReportType,
    ReportVehicleRankingItem,
)


UTC_PLUS_8 = timezone(timedelta(hours=8))
WEIGHT_ZERO = Decimal("0.000")
MONEY_ZERO = Decimal("0.00")


class ReportService:
    """Aggregate completed operations within explicit UTC+8 periods."""

    def __init__(
        self,
        session: Session,
        *,
        now_factory: Callable[[], datetime] = utc_now,
    ) -> None:
        self._session = session
        self._now_factory = now_factory

    def daily(self, report_date: date | None = None) -> BusinessReport:
        selected = report_date or self._local_today()
        return self._build_report(
            report_type="daily",
            period_start=selected,
            period_end_exclusive=selected + timedelta(days=1),
        )

    def monthly(
        self,
        *,
        year: int | None = None,
        month: int | None = None,
    ) -> BusinessReport:
        if (year is None) != (month is None):
            raise BusinessRuleError("year and month must be provided together")
        if year is None or month is None:
            today = self._local_today()
            year, month = today.year, today.month
        if not 1 <= month <= 12:
            raise BusinessRuleError("month must be between 1 and 12")
        period_start = date(year, month, 1)
        days = monthrange(year, month)[1]
        return self._build_report(
            report_type="monthly",
            period_start=period_start,
            period_end_exclusive=period_start + timedelta(days=days),
        )

    def export_xlsx(self, report: BusinessReport) -> bytes:
        """Serialize one already-aggregated report into a three-sheet workbook."""
        workbook = Workbook()
        summary = workbook.active
        summary.title = "经营概览"
        summary.append(("指标", "数值", "单位"))
        summary_rows = (
            ("报表类型", "日报" if report.report_type == "daily" else "月报", ""),
            ("开始日期", report.period_start.isoformat(), ""),
            ("结束日期", report.period_end.isoformat(), ""),
            ("任务数量", report.task_count, "单"),
            ("完成数量", report.completed_task_count, "单"),
            ("完成净重量", report.completed_net_weight_tons, "吨"),
            ("收入", report.income, "元"),
        )
        for row in summary_rows:
            summary.append(row)
        self._style_sheet(summary, widths=(20, 22, 12))
        summary["B7"].number_format = "0.000"
        summary["B8"].number_format = "0.00"

        cargo_sheet = workbook.create_sheet("货物排行")
        cargo_sheet.append(("排名", "货物类型", "完成任务数", "完成净重量(t)"))
        for index, item in enumerate(report.cargo_ranking, start=1):
            cargo_sheet.append(
                (
                    index,
                    item.cargo_type.value,
                    item.completed_task_count,
                    item.total_net_weight_tons,
                )
            )
        self._style_sheet(cargo_sheet, widths=(10, 18, 16, 20))
        for cell in cargo_sheet["D"][1:]:
            cell.number_format = "0.000"

        vehicle_sheet = workbook.create_sheet("车辆排行")
        vehicle_sheet.append(("排名", "车牌号", "完成任务数", "完成净重量(t)"))
        for index, item in enumerate(report.vehicle_ranking, start=1):
            vehicle_sheet.append(
                (
                    index,
                    item.plate_number,
                    item.completed_task_count,
                    item.total_net_weight_tons,
                )
            )
        self._style_sheet(vehicle_sheet, widths=(10, 18, 16, 20))
        for cell in vehicle_sheet["D"][1:]:
            cell.number_format = "0.000"

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _build_report(
        self,
        *,
        report_type: ReportType,
        period_start: date,
        period_end_exclusive: date,
    ) -> BusinessReport:
        start_utc = self._local_day_start_utc(period_start)
        end_utc = self._local_day_start_utc(period_end_exclusive)
        active_task = WeighingTask.deleted_at.is_(None)
        completed_period = (
            active_task,
            WeighingTask.status == WeighingStatus.COMPLETED,
            WeighingTask.completed_at >= start_utc,
            WeighingTask.completed_at < end_utc,
        )
        task_count = self._session.scalar(
            select(func.count(WeighingTask.id)).where(
                active_task,
                WeighingTask.created_at >= start_utc,
                WeighingTask.created_at < end_utc,
            )
        ) or 0
        completed_count = self._session.scalar(
            select(func.count(WeighingTask.id)).where(*completed_period)
        ) or 0
        completed_weight = self._decimal_or(
            self._session.scalar(
                select(func.sum(WeighingTask.net_weight_tons)).where(
                    *completed_period
                )
            ),
            WEIGHT_ZERO,
        )
        income = self._decimal_or(
            self._session.scalar(
                select(func.sum(BillingRecord.fee_amount))
                .join(
                    WeighingTask,
                    WeighingTask.id == BillingRecord.weighing_task_id,
                )
                .where(
                    *completed_period,
                    BillingRecord.payment_status != PaymentStatus.WAIVED,
                )
            ),
            MONEY_ZERO,
        )
        return BusinessReport(
            report_type=report_type,
            period_start=period_start,
            period_end=period_end_exclusive - timedelta(days=1),
            task_count=int(task_count),
            completed_task_count=int(completed_count),
            completed_net_weight_tons=completed_weight.quantize(
                Decimal("0.001")
            ),
            income=income.quantize(Decimal("0.01")),
            cargo_ranking=self._cargo_ranking(completed_period),
            vehicle_ranking=self._vehicle_ranking(completed_period),
        )

    def _cargo_ranking(
        self,
        completed_period: tuple[object, ...],
    ) -> list[ReportCargoRankingItem]:
        total_weight = func.sum(WeighingTask.net_weight_tons)
        rows = self._session.execute(
            select(
                WeighingTask.cargo_type,
                func.count(WeighingTask.id),
                total_weight,
            )
            .where(*completed_period)
            .group_by(WeighingTask.cargo_type)
            .order_by(total_weight.desc(), WeighingTask.cargo_type)
        ).all()
        return [
            ReportCargoRankingItem(
                cargo_type=CargoType(cargo_type),
                completed_task_count=int(task_count),
                total_net_weight_tons=self._decimal_or(
                    weight,
                    WEIGHT_ZERO,
                ).quantize(Decimal("0.001")),
            )
            for cargo_type, task_count, weight in rows
        ]

    def _vehicle_ranking(
        self,
        completed_period: tuple[object, ...],
    ) -> list[ReportVehicleRankingItem]:
        total_weight = func.sum(WeighingTask.net_weight_tons)
        task_count = func.count(WeighingTask.id)
        rows = self._session.execute(
            select(Vehicle.id, Vehicle.plate_number, task_count, total_weight)
            .join(WeighingTask, WeighingTask.vehicle_id == Vehicle.id)
            .where(*completed_period)
            .group_by(Vehicle.id, Vehicle.plate_number)
            .order_by(
                total_weight.desc(),
                task_count.desc(),
                Vehicle.plate_number,
            )
        ).all()
        return [
            ReportVehicleRankingItem(
                vehicle_id=vehicle_id,
                plate_number=plate_number,
                completed_task_count=int(completed_count),
                total_net_weight_tons=self._decimal_or(
                    weight,
                    WEIGHT_ZERO,
                ).quantize(Decimal("0.001")),
            )
            for vehicle_id, plate_number, completed_count, weight in rows
        ]

    def _local_today(self) -> date:
        now = self._now_factory()
        if now.tzinfo is None or now.utcoffset() is None:
            raise ValueError("Report clock must return a timezone-aware datetime")
        return now.astimezone(UTC_PLUS_8).date()

    @staticmethod
    def _local_day_start_utc(value: date) -> datetime:
        return datetime.combine(value, time.min, tzinfo=UTC_PLUS_8).astimezone(
            timezone.utc
        )

    @staticmethod
    def _decimal_or(value: Decimal | int | None, default: Decimal) -> Decimal:
        if value is None:
            return default
        return value if isinstance(value, Decimal) else Decimal(value)

    @staticmethod
    def _style_sheet(worksheet: object, *, widths: tuple[int, ...]) -> None:
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[
                chr(ord("A") + index - 1)
            ].width = width
